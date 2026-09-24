"""GUI smoke test (run under Xvfb): opens on Main Menu -> scan -> STOP ->
scan -> report -> Excel -> browser selector -> Single Link Check -> Settings ->
More (How to Use / Status Codes / About) -> Exit. Needs a display."""
import os, sys, time, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from unittest import mock
from tkinter import messagebox, filedialog
from helpers import Site
import main as M

HTML = lambda b: (lambda h, m: h.reply(200, ("<html>%s</html>" % b).encode()))
routes = {"/": HTML('<a href="/p0">s</a><a href="/gone">x</a><a href="/redir">r</a>'),
          "/redir": lambda h, m: h.reply(301, b"", headers={"Location": "/p0"})}
for i in range(60):
    nxt = '<a href="/p%d">n</a>' % (i + 1) if i < 59 else ""
    routes["/p%d" % i] = HTML('<a href="/dead%d">d</a><img src="/i%d.png">%s' % (i, i, nxt))
site = Site(routes, delay=0.05)

app = M.WebsiteCrawlerApp(settings_path=os.path.join(tempfile.mkdtemp(), "settings.json"))
app.geometry("1000x700")
def pump(sec):
    end = time.time() + sec
    while time.time() < end:
        app.update(); time.sleep(0.01)
pump(0.3)
_stack = app.tk.splitlist(app.tk.call("winfo", "children", app.menu_frame.master))
assert _stack[-1] == str(app.menu_frame), "app must open on the Main Menu"
print("opens on main menu ok"); app.show_menu(); pump(0.2); print("menu ok")
app.show_scanner(); pump(0.2)
sc = app.scanner_frame

# --- scan 1: STOP
sc.url_entry.delete(0, "end"); sc.url_entry.insert(0, site.url("/"))
sc.depth_var.set("Unlimited")
sc.start_scan()
worst = 0; t0 = time.time(); last = time.time()
while time.time() - t0 < 1.5:
    app.update(); now = time.time(); worst = max(worst, now - last); last = now; time.sleep(0.005)
print("mid-scan stats:", sc.stats_label.cget("text").replace("\n", " | "))
print("now:", sc.now_label.cget("text")[:80], "| worst UI gap %.3fs" % worst)
ts = time.time(); sc.stop_scan()
while sc.scanning and time.time() - ts < 10: app.update(); time.sleep(0.01)
print("stopped in %.2fs, status=%s" % (time.time() - ts, sc.status_label.cget("text")))
assert not sc.scanning and "stopped" in sc.status_label.cget("text")

# --- scan 2: complete -> report
sc.depth_var.set("5"); sc.start_scan()
shown = {}
with mock.patch.object(messagebox, "showinfo", lambda *a, **k: shown.setdefault("info", a)):
    ts = time.time()
    while time.time() - ts < 60:
        app.update(); time.sleep(0.01)
        if sc.finished and hasattr(app.report_frame, "counts"): break
    pump(1.0)
    rf = app.report_frame
    assert hasattr(rf, "counts"), "report not loaded"
    print("report shown after %.1fs" % (time.time() - ts))
    print("summary:", [(k, v) for k, v in rf.summary_items()][:12])
    out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
    with mock.patch.object(filedialog, "asksaveasfilename", lambda **k: out):
        rf.download_report()
    import openpyxl
    wb = openpyxl.load_workbook(out)
    for ws in wb: print("sheet", ws.title, ws.max_row, "rows x", ws.max_column, "cols")
    print("excel saved:", bool(shown))

# --- browser selector: values, missing browser dialog, real Chromium scan
app.show_scanner(); pump(0.2)
assert list(sc.mode_combo.cget("values")) == ["Auto", "HTTP", "Chromium", "Chrome",
    "Firefox", "Edge", "Brave", "Opera"], sc.mode_combo.cget("values")
assert sc.mode_var.get() == "HTTP"
print("selector values ok")

asked = {}
def fake_ask(title, msg, **k):
    asked["title"], asked["msg"] = title, msg
    return False
sc.mode_var.set("Opera"); sc.url_entry.delete(0, "end"); sc.url_entry.insert(0, site.url("/"))
with mock.patch.object(messagebox, "askyesno", fake_ask):
    sc.start_scan()
print("missing-browser dialog:", asked.get("title"), "|", asked.get("msg", "").splitlines()[0])
assert not sc.scanning and "Opera" in asked["title"]

sc.mode_var.set("Opera"); sc.depth_var.set("1")
with mock.patch.object(messagebox, "askyesno", lambda *a, **k: True):
    sc.start_scan()                       # user accepts HTTP fallback
assert sc.scanning and sc.crawl_mode_label == "HTTP"
ts = time.time()
with mock.patch.object(messagebox, "showinfo", lambda *a, **k: None):
    while time.time() - ts < 30 and not sc.finished: app.update(); time.sleep(0.01)
print("fallback scan finished:", sc.finished)
assert sc.finished

app.show_scanner(); pump(0.2)
from browser_crawler import check_browser
if check_browser("chromium").ok:
    sc.mode_var.set("Chromium"); sc.depth_var.set("1"); sc.start_scan()
    assert sc.scanning
    ts = time.time()
    while time.time() - ts < 90 and not sc.finished: app.update(); time.sleep(0.01)
    print("chromium scan finished:", sc.finished, "| now:", sc.now_label.cget("text")[:60],
          "| info:", app.report_frame.scan_info)
    assert sc.finished and "Chromium" in app.report_frame.scan_info["Crawl mode"]

# --- menu entries, Single Link Check, Settings, More, About, Exit
app.show_menu(); pump(0.2)
assert list(app.menu_frame.buttons) == ["Full Website Scan", "Single Link Check", "Settings", "More", "Exit"]
app.menu_frame.buttons["Single Link Check"].invoke(); pump(0.2)
sl = app.single_frame
sl.url_entry.insert(0, site.url("/redir")); sl.check_button.invoke()
ts = time.time()
while sl.checking and time.time() - ts < 20: app.update(); time.sleep(0.01)
vals = {k: v.get("1.0", "end").strip() for k, v in sl.values.items()}
print("single check:", vals["HTTP Status"], "|", vals["Status Category"], "|", vals["Redirects"][:60])
assert vals["Status Category"] == "Redirect" and vals["HTTP Status"].startswith("200")
sl.url_entry.delete(0, "end"); sl.url_entry.insert(0, "not a url"); sl.check_button.invoke()
print("invalid url ->", sl.status_label.cget("text"))

app.show_menu(); app.menu_frame.buttons["Settings"].invoke(); pump(0.2)
st = app.settings_frame
st.timeout_var.set("20"); st.redirects_var.set(False); st.save_button.invoke(); pump(0.1)
cur = app.settings_store.current
print("settings saved:", cur)
assert (cur.request_timeout, cur.follow_redirects, cur.respect_robots) == (20, False, True)
sl.url_entry.delete(0, "end"); sl.url_entry.insert(0, site.url("/redir")); sl.check_button.invoke()
while sl.checking: app.update(); time.sleep(0.01)
print("single check, redirects off:", sl.values["HTTP Status"].get("1.0", "end").strip())
assert sl.values["HTTP Status"].get("1.0", "end").startswith("301")

app.show_menu(); app.menu_frame.buttons["More"].invoke(); pump(0.2)
for label in ("How to Use", "HTTP Status Codes", "About"):
    app.show_more(); app.more_frame.buttons[label].invoke(); pump(0.2)
    print("more ->", label, "ok")
with mock.patch("more_screens.webbrowser.open", return_value=True) as opened:
    app.about_frame.paypal_button.invoke()
assert opened.called and "bartoszsartoszek%40gmail.com" in opened.call_args[0][0]
print("paypal opens:", opened.call_args[0][0])

app.show_menu(); pump(0.1)
app.menu_frame.buttons["Exit"].invoke()
try:
    app.winfo_exists(); raise SystemExit("Exit did not close the app")
except Exception as error:
    if isinstance(error, SystemExit): raise
    print("exit closed the app")
print("SMOKE OK")
