import re
import tkinter as tk
from collections import Counter
from datetime import datetime
from tkinter import ttk, filedialog, messagebox

from models import (
    CAT_BROKEN,
    CAT_OK,
    CAT_REDIRECT,
    CAT_UNVERIFIED,
    describe_result,
)
from utils import format_time

_CATEGORY_LABEL = {
    CAT_OK: "OK",
    CAT_REDIRECT: "Redirect",
    CAT_BROKEN: "Broken",
    CAT_UNVERIFIED: "Unverified",
}
# Problems tab / sheet order: definitely-dead links first.
_PROBLEM_ORDER = {CAT_BROKEN: 0, CAT_UNVERIFIED: 1}
_UNVERIFIED_KINDS = ("Timeout", "Server Error", "SSL/Connection", "Blocked/Unverified")

_ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _excel_safe(value):
    """Keep openpyxl from crashing on control characters or treating text
    that starts with '=' as a formula."""
    if isinstance(value, str):
        value = _ILLEGAL_XLSX_CHARS.sub("", value)
        if value.startswith("="):
            value = "'" + value
    return value


class ScanReportFrame(tk.Frame):
    TABLE_COLUMNS = (
        ("status", "Status", 120, "center", False),
        ("url", "URL", 280, "w", False),
        ("source_page", "Source Page", 260, "w", False),
        ("source_type", "Source Type", 110, "w", False),
        ("link_text", "Link Text", 180, "w", False),
        ("final_url", "Final URL", 280, "w", False),
        ("response_time", "Response Time", 115, "center", False),
        ("details", "Details", 320, "w", False),
    )
    ROWS_PER_CHUNK = 300
    URL_COLUMN = 1     # the "URL" column: the link itself, not its source page

    def __init__(self, parent, back_to_scanner):
        super().__init__(parent, bg="#1e1e1e")
        self.back_to_scanner_callback = back_to_scanner
        self.website_url = ""
        self.elapsed = 0
        self.pages_scanned = 0
        self.results = []
        self.scan_info = {}
        self.rows = []
        self.counts = {}
        self.url_context_menu = None
        self.active_tree = None
        self.copy_bars = {}       # tree -> (copy button, info label)
        self._generation = 0

    def load_report(self, website_url, elapsed, pages_scanned, results,
                    scan_info=None):
        self.website_url = website_url
        self.elapsed = elapsed
        self.pages_scanned = pages_scanned
        self.results = list(results)
        self.scan_info = dict(scan_info or {})
        self._generation += 1

        # One classification pass shared by the GUI and the Excel export.
        self.rows = [
            (link, result, describe_result(link, result))
            for link, result in self.results
        ]
        self.counts = self.compute_counts()

        for widget in self.winfo_children():
            widget.destroy()

        self.copy_bars = {}
        self.create_url_context_menu()
        self.build_ui()

    # ---------------------------------------------------------
    # Numbers (single source of truth for GUI + Excel)
    # ---------------------------------------------------------

    def compute_counts(self):
        categories = Counter(row["category"] for _l, _r, row in self.rows)
        kinds = Counter(
            row["problem_kind"] for _l, _r, row in self.rows
            if row["category"] == CAT_UNVERIFIED
        )
        return {
            "ok": categories[CAT_OK],
            "redirect": categories[CAT_REDIRECT],
            "broken": categories[CAT_BROKEN],
            "unverified": categories[CAT_UNVERIFIED],
            "kinds": kinds,
        }

    def count_results(self, category):
        return self.counts.get(category, 0)

    def summary_items(self):
        items = [("Website", self.website_url)]
        for key, value in self.scan_info.items():
            items.append((key, str(value)))
        total = len(self.rows)
        items += [
            ("Pages scanned", str(self.pages_scanned)),
            ("Links discovered", str(total)),
            ("Links checked", str(total)),
            ("OK", str(self.counts["ok"])),
            ("Redirects", str(self.counts["redirect"])),
            ("Broken", str(self.counts["broken"])),
            ("Unverified", str(self.counts["unverified"])),
        ]
        for kind in _UNVERIFIED_KINDS:
            if self.counts["kinds"].get(kind):
                items.append((f"   ↳ {kind}", str(self.counts["kinds"][kind])))
        items.append(("Scan time", self.format_time(self.elapsed)))
        return items

    @staticmethod
    def requires_attention(result):
        return result.is_broken or result.is_unverified

    @staticmethod
    def get_result_display_data(link, result):
        """Kept for compatibility; delegates to the shared describer."""
        row = describe_result(link, result)
        return (row["status"], row["final_url"], row["response_time"],
                row["details"], row["tag"])

    # ---------------------------------------------------------
    # UI
    # ---------------------------------------------------------

    def create_url_context_menu(self):
        self.url_context_menu = tk.Menu(self, tearoff=False)
        self.url_context_menu.add_command(
            label="Copy URL",
            command=self.copy_selected_url
        )

    def build_ui(self):
        title = tk.Label(
            self,
            text="Scan Report",
            font=("Arial", 30, "bold"),
            fg="white",
            bg="#1e1e1e"
        )
        title.pack(pady=(20, 5))

        completed = tk.Label(
            self,
            text="SCAN COMPLETED",
            font=("Arial", 14, "bold"),
            fg="#4caf50",
            bg="#1e1e1e"
        )
        completed.pack(pady=(0, 10))

        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True, padx=35, pady=10)

        summary_tab = tk.Frame(notebook, bg="#1e1e1e")
        problems_tab = tk.Frame(notebook, bg="#1e1e1e")
        redirects_tab = tk.Frame(notebook, bg="#1e1e1e")
        all_tab = tk.Frame(notebook, bg="#1e1e1e")

        notebook.add(summary_tab, text="Summary")
        notebook.add(problems_tab, text=f"Problems ({self.counts['broken'] + self.counts['unverified']})")
        notebook.add(redirects_tab, text=f"Redirects ({self.counts['redirect']})")
        notebook.add(all_tab, text=f"All Links ({len(self.rows)})")

        self.build_summary(summary_tab)

        problem_rows = sorted(
            (item for item in self.rows if item[2]["category"] in _PROBLEM_ORDER),
            key=lambda item: _PROBLEM_ORDER[item[2]["category"]],
        )
        redirect_rows = [
            item for item in self.rows if item[2]["category"] == CAT_REDIRECT
        ]

        self.create_table(problems_tab, problem_rows,
                          empty_text="✓ No problems found", empty_ok=True)
        self.create_table(redirects_tab, redirect_rows,
                          empty_text="No redirects found")
        self.create_table(all_tab, self.rows, empty_text="No links found")

        bottom = tk.Frame(self, bg="#1e1e1e")
        bottom.pack(fill="x", padx=35, pady=(0, 20))

        tk.Button(
            bottom,
            text="← BACK TO SCANNER",
            font=("Arial", 12, "bold"),
            bg="#444444",
            fg="white",
            width=20,
            height=2,
            command=self.back_to_scanner
        ).pack(side="left")

        tk.Button(
            bottom,
            text="DOWNLOAD REPORT",
            font=("Arial", 12, "bold"),
            bg="#2d8cf0",
            fg="white",
            width=20,
            height=2,
            command=self.download_report
        ).pack(side="right")

    def build_summary(self, parent):
        summary = tk.Frame(parent, bg="#252525")
        summary.pack(fill="x", padx=15, pady=15)

        for row, (label, value) in enumerate(self.summary_items()):
            tk.Label(
                summary,
                text=f"{label}:",
                font=("Arial", 11, "bold"),
                fg="white",
                bg="#252525",
                anchor="w"
            ).grid(row=row, column=0, sticky="w", padx=15, pady=4)

            tk.Label(
                summary,
                text=value,
                font=("Arial", 11),
                fg="white",
                bg="#252525",
                anchor="w"
            ).grid(row=row, column=1, sticky="w", padx=15, pady=4)

    def create_table(self, parent, rows, empty_text, empty_ok=False):
        # Grid layout: table + vertical bar on top, horizontal bar underneath
        # (always visible), then the Copy URL bar.
        frame = tk.Frame(parent, bg="#1e1e1e")
        frame.pack(fill="both", expand=True, padx=10, pady=(10, 4))
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)

        columns = tuple(col[0] for col in self.TABLE_COLUMNS)
        tree = ttk.Treeview(frame, columns=columns, show="headings")

        for key, heading, width, anchor, stretch in self.TABLE_COLUMNS:
            tree.heading(key, text=heading)
            tree.column(key, width=width, minwidth=90, anchor=anchor,
                        stretch=stretch)

        tree.tag_configure("ok", foreground="green")
        tree.tag_configure("redirect", foreground="orange")
        tree.tag_configure("error", foreground="red")

        scrollbar = ttk.Scrollbar(
            frame,
            orient="vertical",
            command=tree.yview
        )
        horizontal_scrollbar = ttk.Scrollbar(
            frame,
            orient="horizontal",
            command=tree.xview
        )
        tree.configure(
            yscrollcommand=scrollbar.set,
            xscrollcommand=horizontal_scrollbar.set
        )
        tree.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")
        horizontal_scrollbar.grid(row=1, column=0, sticky="ew")

        tree.bind("<Control-c>", self.copy_selected_url)
        tree.bind("<Button-3>", self.show_url_context_menu)
        tree.bind("<<TreeviewSelect>>", self.update_copy_bar)
        # Shift + mouse wheel scrolls sideways (plain wheel stays vertical).
        for sequence in ("<Shift-MouseWheel>", "<Shift-Button-4>",
                         "<Shift-Button-5>"):
            tree.bind(sequence, self.scroll_sideways)

        self.build_copy_bar(parent, tree)

        if not rows:
            tk.Label(
                parent,
                text=empty_text,
                font=("Arial", 16, "bold"),
                fg="#4caf50" if empty_ok else "white",
                bg="#1e1e1e",
            ).place(relx=0.5, rely=0.5, anchor="center")
            return

        # Large scans can produce tens of thousands of rows; inserting them
        # in small chunks keeps the window responsive while it fills.
        self.fill_tree(tree, rows, 0, self._generation)

    def fill_tree(self, tree, rows, start, generation):
        if generation != self._generation:
            return
        end = min(start + self.ROWS_PER_CHUNK, len(rows))
        try:
            for _link, _result, row in rows[start:end]:
                tree.insert(
                    "",
                    "end",
                    values=(
                        row["status"],
                        row["url"],
                        row["source_page"],
                        row["source_type"],
                        row["link_text"],
                        row["final_url"],
                        row["response_time"],
                        row["details"],
                    ),
                    tags=(row["tag"],)
                )
            if end < len(rows):
                self.after(
                    5, lambda: self.fill_tree(tree, rows, end, generation))
        except tk.TclError:
            pass  # the report was closed/rebuilt while filling

    def build_copy_bar(self, parent, tree):
        """Under every table: shows the selected URL and offers Copy URL."""
        bar = tk.Frame(parent, bg="#1e1e1e")
        bar.pack(fill="x", padx=10, pady=(0, 8))
        button = tk.Button(
            bar, text="Copy URL", font=("Arial", 10, "bold"), width=12,
            bg="#2d8cf0", fg="white", state="disabled",
            command=lambda: self.copy_selected_url(tree=tree))
        button.pack(side="left")
        info = tk.Label(
            bar, text="Click a row to select it; right-click for Copy URL.",
            font=("Arial", 10), fg="#aaaaaa", bg="#1e1e1e", anchor="w")
        info.pack(side="left", padx=10, fill="x", expand=True)
        self.copy_bars[tree] = (button, info)

    def selected_url(self, tree):
        """URL of the selected result (the link itself, never its source
        page), or None."""
        selected_items = tree.selection()
        if not selected_items:
            return None
        values = tree.item(selected_items[0], "values")
        return str(values[self.URL_COLUMN]) if values else None

    def update_copy_bar(self, event):
        tree = event.widget
        bar = self.copy_bars.get(tree)
        if bar is None:
            return
        button, info = bar
        url = self.selected_url(tree)
        button.config(state="normal" if url else "disabled")
        info.config(text=f"Selected: {url}" if url else
                    "Click a row to select it; right-click for Copy URL.")

    @staticmethod
    def scroll_sideways(event):
        left = getattr(event, "num", 0) == 4 or getattr(event, "delta", 0) > 0
        event.widget.xview_scroll(-3 if left else 3, "units")
        return "break"

    def show_url_context_menu(self, event):
        tree = event.widget
        item = tree.identify_row(event.y)

        if not item:
            return

        tree.selection_set(item)
        tree.focus(item)
        self.active_tree = tree

        try:
            self.url_context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.url_context_menu.grab_release()

    def copy_selected_url(self, event=None, tree=None):
        """Copy the selected result's URL (Copy URL button, context menu and
        Ctrl+C all end up here)."""
        if tree is None:
            tree = event.widget if event else self.active_tree
        if tree is None:
            return "break"

        url = self.selected_url(tree)
        if not url:
            return "break"

        self.clipboard_clear()
        self.clipboard_append(url)
        self.update_idletasks()
        bar = self.copy_bars.get(tree)
        if bar is not None:
            bar[1].config(text=f"Copied: {url}")
        return "break"

    # ---------------------------------------------------------
    # Excel
    # ---------------------------------------------------------

    EXCEL_HEADERS = [
        "Category",
        "Problem Type",
        "Status",
        "URL",
        "Source Page",
        "Source Type",
        "Source Link Text",
        "Final URL",
        "Response Time (s)",
        "Method",
        "Depth",
        "Occurrences",
        "Details",
    ]
    EXCEL_WIDTHS = [12, 18, 25, 75, 75, 20, 35, 75, 18, 10, 8, 12, 90]

    @staticmethod
    def excel_row(row):
        return [
            _CATEGORY_LABEL[row["category"]],
            row["problem_kind"],
            row["status"],
            row["url"],
            row["source_page"],
            row["source_type"],
            row["link_text"],
            row["final_url"],
            row["response_seconds"],
            row["method"],
            row["depth"],
            row["occurrences"],
            row["details"],
        ]

    def build_workbook(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"

        # Same numbers, in the same order, as the Summary tab in the GUI.
        for label, value in self.summary_items():
            summary_sheet.append([_excel_safe(label.strip()), _excel_safe(value)])
        summary_sheet.append(
            ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        for cell in summary_sheet["A"]:
            cell.font = Font(bold=True)
        summary_sheet.column_dimensions["A"].width = 25
        summary_sheet.column_dimensions["B"].width = 80

        problems_sheet = workbook.create_sheet("Problems")
        redirects_sheet = workbook.create_sheet("Redirects")
        all_sheet = workbook.create_sheet("All Links")
        occurrences_sheet = workbook.create_sheet("Occurrences")

        for sheet in (problems_sheet, redirects_sheet, all_sheet):
            sheet.append(self.EXCEL_HEADERS)

        occurrences_sheet.append([
            "Category", "Status", "URL", "Found On (Source Page)",
            "Source Type", "Link Text",
        ])

        problem_items = sorted(
            (item for item in self.rows if item[2]["category"] in _PROBLEM_ORDER),
            key=lambda item: _PROBLEM_ORDER[item[2]["category"]],
        )

        for _link, _result, row in self.rows:
            data = [_excel_safe(v) for v in self.excel_row(row)]
            all_sheet.append(data)
            if row["category"] == CAT_REDIRECT:
                redirects_sheet.append(data)

        for link, _result, row in problem_items:
            problems_sheet.append(
                [_excel_safe(v) for v in self.excel_row(row)])
            # Every page a broken/unverified URL was found on.
            for source_page, source_type, text in (
                link.occurrences or [(link.source_url, link.source_type,
                                      link.source_text)]
            ):
                occurrences_sheet.append([
                    _excel_safe(_CATEGORY_LABEL[row["category"]]),
                    _excel_safe(row["status"]),
                    _excel_safe(row["url"]),
                    _excel_safe(source_page),
                    _excel_safe(source_type),
                    _excel_safe(text),
                ])

        if not problem_items:
            problems_sheet.append(["✓ No problems found"])

        for sheet in (problems_sheet, redirects_sheet, all_sheet):
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center")
            for index, width in enumerate(self.EXCEL_WIDTHS, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            sheet.freeze_panes = "A2"

        for cell in occurrences_sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")
        for index, width in enumerate([12, 25, 75, 75, 20, 40], start=1):
            occurrences_sheet.column_dimensions[
                get_column_letter(index)].width = width
        occurrences_sheet.freeze_panes = "A2"

        return workbook

    def download_report(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            messagebox.showerror(
                "Missing package",
                "Install openpyxl with:\n\npython3 -m pip install openpyxl"
            )
            return

        filename = filedialog.asksaveasfilename(
            title="Save scan report",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="website_scan_report.xlsx"
        )

        if not filename:
            return

        try:
            self.build_workbook().save(filename)
        except OSError as error:
            messagebox.showerror(
                "Could not save report",
                f"{filename}\n\n{error}"
            )
            return

        messagebox.showinfo(
            "Report saved",
            f"Report saved successfully:\n\n{filename}"
        )

    def back_to_scanner(self):
        self.back_to_scanner_callback()

    @staticmethod
    def format_time(seconds):
        return format_time(seconds)
